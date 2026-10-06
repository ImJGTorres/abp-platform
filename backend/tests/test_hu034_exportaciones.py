"""
HU-034 — Exportar reportes en PDF y Excel.
"""
from unittest.mock import patch, MagicMock
from django.test import RequestFactory

from apps.exportaciones.views import SolicitarExportacionView, EstadoExportacionView
from tests.conftest_hu import make_payload, authenticated_request, auth_ctx


class TestSolicitarExportacionView:

    def setup_method(self):
        self.factory = RequestFactory()

    @patch('apps.exportaciones.views.generar_exportacion')
    @patch('apps.exportaciones.models.Exportacion.objects.create')
    def test_solicitar_pdf_retorna_201_con_id_y_estado(self, mock_create, mock_gen):
        mock_exp = MagicMock()
        mock_exp.id = 42
        mock_exp.estado = 'listo'
        mock_exp.mensaje_error = None
        mock_create.return_value = mock_exp

        req = authenticated_request(
            self.factory, 'POST', '/api/exportar/reporte/',
            make_payload(tipo_rol='director'),
            data={'tipo_reporte': 'proyecto', 'formato': 'pdf',
                  'parametros': {'proyecto_id': 1}},
        )
        with auth_ctx(req):
            response = SolicitarExportacionView.as_view()(req)

        assert response.status_code == 201
        assert 'id' in response.data
        assert 'estado' in response.data

    @patch('apps.exportaciones.views.generar_exportacion')
    @patch('apps.exportaciones.models.Exportacion.objects.create')
    def test_solicitar_excel_retorna_201(self, mock_create, mock_gen):
        mock_exp = MagicMock()
        mock_exp.id = 43
        mock_exp.estado = 'listo'
        mock_exp.mensaje_error = None
        mock_create.return_value = mock_exp

        req = authenticated_request(
            self.factory, 'POST', '/api/exportar/reporte/',
            make_payload(tipo_rol='director'),
            data={'tipo_reporte': 'proyecto', 'formato': 'excel',
                  'parametros': {'proyecto_id': 1}},
        )
        with auth_ctx(req):
            response = SolicitarExportacionView.as_view()(req)

        assert response.status_code == 201

    @patch('apps.exportaciones.models.Exportacion.objects.create')
    def test_formato_no_permitido_retorna_400_sin_crear_registro(self, mock_create):
        req = authenticated_request(
            self.factory, 'POST', '/api/exportar/reporte/',
            make_payload(tipo_rol='director'),
            data={'tipo_reporte': 'proyecto', 'formato': 'word',
                  'parametros': {'proyecto_id': 1}},
        )
        with auth_ctx(req):
            response = SolicitarExportacionView.as_view()(req)

        assert response.status_code == 400
        mock_create.assert_not_called()

    @patch('apps.exportaciones.models.Exportacion.objects.create')
    def test_tipo_reporte_desconocido_retorna_400(self, mock_create):
        req = authenticated_request(
            self.factory, 'POST', '/api/exportar/reporte/',
            make_payload(tipo_rol='director'),
            data={'tipo_reporte': 'desconocido', 'formato': 'pdf', 'parametros': {}},
        )
        with auth_ctx(req):
            response = SolicitarExportacionView.as_view()(req)

        assert response.status_code == 400

    @patch('apps.exportaciones.models.Exportacion.objects.create')
    def test_falta_proyecto_id_para_tipo_proyecto_retorna_400(self, mock_create):
        req = authenticated_request(
            self.factory, 'POST', '/api/exportar/reporte/',
            make_payload(tipo_rol='director'),
            data={'tipo_reporte': 'proyecto', 'formato': 'pdf', 'parametros': {}},
        )
        with auth_ctx(req):
            response = SolicitarExportacionView.as_view()(req)

        assert response.status_code == 400


class TestEstadoExportacionView:

    def setup_method(self):
        self.factory = RequestFactory()

    @patch('apps.exportaciones.models.Exportacion.objects.get')
    def test_exportacion_lista_incluye_url_descarga(self, mock_get):
        mock_exp = MagicMock()
        mock_exp.id = 1
        mock_exp.tipo_reporte = 'proyecto'
        mock_exp.formato = 'pdf'
        mock_exp.estado = 'listo'
        mock_exp.ruta_archivo = '/fake/media/exportaciones/abc.pdf'
        mock_exp.fecha_solicitud.isoformat.return_value = '2026-05-01T00:00:00'
        mock_exp.fecha_disponible.isoformat.return_value = '2026-05-01T00:01:00'
        mock_exp.mensaje_error = None
        mock_get.return_value = mock_exp

        req = authenticated_request(self.factory, 'GET', '/api/exportar/1/estado/',
                                    make_payload(user_id=1, tipo_rol='director'))
        with auth_ctx(req), patch('apps.exportaciones.views.os.path.relpath',
                                  return_value='exportaciones/abc.pdf'):
            response = EstadoExportacionView.as_view()(req, exportacion_id=1)

        assert response.status_code == 200
        assert response.data['estado'] == 'listo'
        assert response.data['url_descarga'] is not None

    @patch('apps.exportaciones.models.Exportacion.objects.get')
    def test_exportacion_generando_no_tiene_url_descarga(self, mock_get):
        mock_exp = MagicMock()
        mock_exp.id = 2
        mock_exp.tipo_reporte = 'proyecto'
        mock_exp.formato = 'pdf'
        mock_exp.estado = 'generando'
        mock_exp.ruta_archivo = None
        mock_exp.fecha_solicitud.isoformat.return_value = '2026-05-01T00:00:00'
        mock_exp.fecha_disponible = None
        mock_exp.mensaje_error = None
        mock_get.return_value = mock_exp

        req = authenticated_request(self.factory, 'GET', '/api/exportar/2/estado/',
                                    make_payload(user_id=1, tipo_rol='director'))
        with auth_ctx(req):
            response = EstadoExportacionView.as_view()(req, exportacion_id=2)

        assert response.status_code == 200
        assert response.data['url_descarga'] is None

    @patch('apps.exportaciones.models.Exportacion.objects.get')
    def test_exportacion_en_error_incluye_mensaje_error(self, mock_get):
        """Estado error → url_descarga nula y mensaje_error presente en respuesta."""
        mock_exp = MagicMock()
        mock_exp.id = 3
        mock_exp.tipo_reporte = 'proyecto'
        mock_exp.formato = 'pdf'
        mock_exp.estado = 'error'
        mock_exp.ruta_archivo = None
        mock_exp.fecha_solicitud.isoformat.return_value = '2026-05-01T00:00:00'
        mock_exp.fecha_disponible = None
        mock_exp.mensaje_error = 'Proyecto no encontrado al generar PDF'
        mock_get.return_value = mock_exp

        req = authenticated_request(self.factory, 'GET', '/api/exportar/3/estado/',
                                    make_payload(user_id=1, tipo_rol='director'))
        with auth_ctx(req):
            response = EstadoExportacionView.as_view()(req, exportacion_id=3)

        assert response.status_code == 200
        assert response.data['estado'] == 'error'
        assert response.data['url_descarga'] is None
        assert response.data['mensaje_error'] == 'Proyecto no encontrado al generar PDF'

    @patch('apps.exportaciones.models.Exportacion.objects.get')
    def test_exportacion_de_otro_usuario_retorna_404(self, mock_get):
        from apps.exportaciones.models import Exportacion
        mock_get.side_effect = Exportacion.DoesNotExist

        req = authenticated_request(self.factory, 'GET', '/api/exportar/99/estado/',
                                    make_payload(user_id=1, tipo_rol='director'))
        with auth_ctx(req):
            response = EstadoExportacionView.as_view()(req, exportacion_id=99)

        assert response.status_code == 404


def test_criterio_rojo_muestra_nota_y_porcentaje_del_semaforo():
    from apps.exportaciones.generadores import _criterio_rojo
    datos = {'umbrales_semaforo': {'nota_rojo': 3.0, 'pct_rojo': 50.0}}
    assert _criterio_rojo(datos) == 'Riesgo crítico (rojo): nota < 3 o actividades incumplidas >= 50%'
    # Datos sin umbrales_semaforo (formato anterior) siguen funcionando
    assert 'nota < 2.5' in _criterio_rojo({'umbral_bajo_rendimiento': 2.5})


# ── SCRUM-602: membrete institucional en los PDF ─────────────────────────────

import pytest
from apps.exportaciones.generadores.pdf import (
    generar_pdf_estudiante, generar_pdf_indicadores, generar_pdf_proyecto,
)

GENERADORES_PDF = [generar_pdf_proyecto, generar_pdf_estudiante, generar_pdf_indicadores]


def _identidad(logo=None):
    return {'nombre_institucion': 'Universidad de Prueba', 'programa_academico': 'Programa de Prueba',
            'logotipo_path': logo, 'logotipo_url': None}


def _pdf(generador, tmp_path, identidad):
    """Genera el PDF sin compresión para poder buscar texto e imágenes en los bytes."""
    ruta = str(tmp_path / f'{generador.__name__}.pdf')
    with patch('apps.configuracion.identidad.obtener_identidad', return_value=identidad), \
         patch('reportlab.rl_config.pageCompression', 0):
        generador({}, ruta)
    return open(ruta, 'rb').read()


@pytest.fixture
def logo_png(tmp_path):
    from PIL import Image
    ruta = tmp_path / 'logo.png'
    Image.new('RGB', (200, 100), 'red').save(ruta)
    return str(ruta)


@pytest.mark.parametrize('generador', GENERADORES_PDF)
def test_pdf_muestra_membrete_y_logo_configurados(generador, tmp_path, logo_png):
    pdf = _pdf(generador, tmp_path, _identidad(logo_png))
    assert b'Universidad de Prueba' in pdf
    assert b'Programa de Prueba' in pdf
    assert b'/Subtype /Image' in pdf


@pytest.mark.parametrize('generador', GENERADORES_PDF)
def test_pdf_sin_logo_muestra_solo_texto(generador, tmp_path):
    pdf = _pdf(generador, tmp_path, _identidad(None))
    assert b'Universidad de Prueba' in pdf
    assert b'/Subtype /Image' not in pdf


@pytest.mark.parametrize('generador', GENERADORES_PDF)
def test_pdf_se_genera_con_logo_borrado_del_disco(generador, tmp_path):
    pdf = _pdf(generador, tmp_path, _identidad(str(tmp_path / 'no_existe.png')))
    assert b'Universidad de Prueba' in pdf
    assert b'/Subtype /Image' not in pdf


@pytest.mark.parametrize('generador', GENERADORES_PDF)
def test_pdf_se_genera_con_logo_danado(generador, tmp_path):
    danado = tmp_path / 'danado.png'
    danado.write_bytes(b'esto no es una imagen')
    pdf = _pdf(generador, tmp_path, _identidad(str(danado)))
    assert b'Universidad de Prueba' in pdf
    assert b'/Subtype /Image' not in pdf


def test_pdf_usa_valores_por_defecto_si_falla_la_identidad(tmp_path):
    ruta = str(tmp_path / 'p.pdf')
    with patch('apps.configuracion.identidad.obtener_identidad', side_effect=RuntimeError('bd caída')), \
         patch('reportlab.rl_config.pageCompression', 0):
        generar_pdf_proyecto({}, ruta)
    assert b'Ingenier' in open(ruta, 'rb').read()  # PROGRAMA por defecto


# ── SCRUM-603: membrete institucional en los Excel ───────────────────────────

from openpyxl import load_workbook
from apps.exportaciones.generadores.excel import (
    generar_excel_estudiante, generar_excel_indicadores, generar_excel_proyecto,
)

# (generador, primera hoja de datos, fila donde queda el encabezado de su tabla, primer encabezado)
GENERADORES_EXCEL = [
    (generar_excel_proyecto, 'Fases', 5, 'Fase'),
    (generar_excel_estudiante, 'Nota Final', 5, 'Componente'),
    (generar_excel_indicadores, 'Resumen', 3, 'Periodo'),
]


def _xlsx(generador, tmp_path, identidad):
    ruta = str(tmp_path / f'{generador.__name__}.xlsx')
    with patch('apps.configuracion.identidad.obtener_identidad', return_value=identidad):
        generador({}, ruta)
    return load_workbook(ruta)


@pytest.mark.parametrize('generador,hoja,fila_tabla,encabezado', GENERADORES_EXCEL)
def test_excel_muestra_membrete_y_logo_configurados(generador, hoja, fila_tabla, encabezado, tmp_path, logo_png):
    wb = _xlsx(generador, tmp_path, _identidad(logo_png))
    info = wb['Info']
    assert len(info._images) == 1
    ancla = info._images[0].anchor  # tamaño mostrado, en EMU (9525 por px)
    assert (ancla._from.col, ancla._from.row) == (0, 0)  # A1
    assert ancla.ext.height / 9525 == 60
    assert info['C1'].value == 'Universidad de Prueba'
    assert info['C2'].value == 'Programa de Prueba'
    assert info['C3'].value.startswith('Exportado:')
    datos = wb[hoja]
    assert (datos['A1'].value, datos['A2'].value) == ('Universidad de Prueba', 'Programa de Prueba')
    assert datos.cell(fila_tabla, 1).value == encabezado  # la tabla bajó 2 filas intacta


@pytest.mark.parametrize('generador,hoja,fila_tabla,encabezado', GENERADORES_EXCEL)
def test_excel_sin_logo_escribe_membrete_en_columna_a(generador, hoja, fila_tabla, encabezado, tmp_path):
    info = _xlsx(generador, tmp_path, _identidad(None))['Info']
    assert info._images == []
    assert (info['A1'].value, info['A2'].value) == ('Universidad de Prueba', 'Programa de Prueba')
    assert info['A3'].value.startswith('Exportado:')


@pytest.mark.parametrize('generador,hoja,fila_tabla,encabezado', GENERADORES_EXCEL)
def test_excel_se_genera_con_logo_borrado_o_danado(generador, hoja, fila_tabla, encabezado, tmp_path):
    danado = tmp_path / 'danado.png'
    danado.write_bytes(b'esto no es una imagen')
    for logo in (str(tmp_path / 'no_existe.png'), str(danado)):
        info = _xlsx(generador, tmp_path, _identidad(logo))['Info']
        assert info._images == []
        assert info['A1'].value == 'Universidad de Prueba'


# ── SCRUM-604: aceptación — cambiar identidad real (endpoint HU-035) y exportar ──
# A diferencia de los tests de arriba (que mockean obtener_identidad), este usa el
# pipeline real de punta a punta: PUT /api/configuracion/identidad/ → generar_exportacion()
# → archivo en disco, sin mocks de identidad.

import io
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APIClient

from apps.exportaciones.models import Exportacion
from apps.exportaciones.services import generar_exportacion
from tests.factories import AdminFactory, ProyectoFactory


def _png_logo(nombre='logo.png', tamano=(60, 60)):
    from PIL import Image as PILImage
    buf = io.BytesIO()
    PILImage.new('RGB', tamano, 'red').save(buf, 'PNG')
    return SimpleUploadedFile(nombre, buf.getvalue(), content_type='image/png')


@pytest.fixture
def media_temporal(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path


@pytest.mark.django_db
class TestAceptacionCambiarIdentidadYExportar:
    """
    Escenario de aceptación SCRUM-604: un administrador cambia la identidad
    institucional vía el endpoint real de HU-035 (PUT /api/configuracion/identidad/,
    SCRUM-607) y, al exportar un reporte en PDF y en Excel con el pipeline real
    (generar_exportacion, SCRUM-519... es decir el de exportaciones), ambos
    archivos muestran el nombre y el logotipo recién configurados.
    """

    def test_pdf_y_excel_reales_usan_la_identidad_recien_configurada(self, media_temporal):
        admin = AdminFactory()
        cliente = APIClient()
        cliente.force_authenticate(user=admin)

        r = cliente.put('/api/configuracion/identidad/', {
            'nombre_institucion': 'Universidad de Aceptacion',
            'programa_academico': 'Programa de Aceptacion',
            'logotipo': _png_logo(),
        }, format='multipart')
        assert r.status_code == 200

        proyecto = ProyectoFactory()
        exp_pdf = Exportacion.objects.create(
            id_usuario=admin, tipo_reporte='proyecto', formato='pdf',
            parametros={'proyecto_id': proyecto.id},
        )
        exp_excel = Exportacion.objects.create(
            id_usuario=admin, tipo_reporte='proyecto', formato='excel',
            parametros={'proyecto_id': proyecto.id},
        )

        with patch('reportlab.rl_config.pageCompression', 0):
            generar_exportacion(exp_pdf.id)
        generar_exportacion(exp_excel.id)

        exp_pdf.refresh_from_db()
        exp_excel.refresh_from_db()
        assert exp_pdf.estado == 'listo', exp_pdf.mensaje_error
        assert exp_excel.estado == 'listo', exp_excel.mensaje_error

        pdf_bytes = open(exp_pdf.ruta_archivo, 'rb').read()
        assert b'Universidad de Aceptacion' in pdf_bytes
        assert b'/Subtype /Image' in pdf_bytes

        wb = load_workbook(exp_excel.ruta_archivo)
        info = wb['Info']
        assert info['C1'].value == 'Universidad de Aceptacion'
        assert len(info._images) == 1
