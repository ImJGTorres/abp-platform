import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROJO = 'FFD32F2F'
ROJO_CLARO = 'FFFFCDD2'
GRIS = 'FFF5F5F5'


def _estilo_encabezado(cell):
    cell.font = Font(bold=True, color='FFFFFFFF', size=10)
    cell.fill = PatternFill(fill_type='solid', fgColor=ROJO)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _estilo_dato(cell, alternado=False):
    cell.fill = PatternFill(fill_type='solid', fgColor=GRIS if alternado else 'FFFFFFFF')
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.font = Font(size=9)


def _autoajustar(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or '')))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def _hoja_info(wb, nombre_institucion="UFPS — Plataforma ABP"):
    ws = wb.active
    ws.title = "Info"
    ws['A1'] = nombre_institucion
    ws['A1'].font = Font(bold=True, size=14, color=ROJO)
    ws['A2'] = f"Exportado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws['A2'].font = Font(size=9, italic=True)
    return ws


def _escribir_tabla(ws, fila_inicio, encabezados, filas):
    for col_idx, enc in enumerate(encabezados, start=1):
        cell = ws.cell(row=fila_inicio, column=col_idx, value=enc)
        _estilo_encabezado(cell)
    for row_idx, fila in enumerate(filas, start=fila_inicio + 1):
        alternado = (row_idx % 2 == 0)
        for col_idx, valor in enumerate(fila, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=valor)
            _estilo_dato(cell, alternado)


def generar_excel_proyecto(datos, ruta_destino):
    wb = Workbook()
    _hoja_info(wb)

    proy = datos.get('proyecto', {})
    ws_info = wb.active
    ws_info['A4'] = 'Proyecto:'
    ws_info['B4'] = proy.get('nombre', '-')
    ws_info['A5'] = 'Estado:'
    ws_info['B5'] = proy.get('estado', '-')
    ws_info['A6'] = 'Curso:'
    ws_info['B6'] = proy.get('curso', '-')

    ws_fases = wb.create_sheet('Fases')
    ws_fases['A1'] = f"Proyecto: {proy.get('nombre','')}"
    ws_fases['A1'].font = Font(bold=True, color=ROJO)
    fases = datos.get('progreso_por_fase', [])
    encabezados = ['Fase', 'Estado', 'Avance (%)', 'Total actividades', 'Completadas', 'En progreso', 'Bloqueadas']
    filas = [
        [f.get('nombre_fase'), f.get('estado_fase'), f.get('porcentaje_almacenado'),
         f.get('total_actividades'), f.get('actividades_completadas'),
         f.get('actividades_en_progreso'), f.get('actividades_bloqueadas')]
        for f in fases
    ]
    _escribir_tabla(ws_fases, 3, encabezados, filas)
    _autoajustar(ws_fases)

    ws_ents = wb.create_sheet('Entregables')
    ws_ents['A1'] = f"Entregables — {proy.get('nombre','')}"
    ws_ents['A1'].font = Font(bold=True, color=ROJO)
    ents = datos.get('resumen_entregables', {})
    _escribir_tabla(ws_ents, 3,
                    ['Indicador', 'Valor'],
                    [
                        ['Total', ents.get('total_entregables', 0)],
                        ['Aprobados', ents.get('aprobados', 0)],
                        ['Rechazados', ents.get('rechazados', 0)],
                        ['Enviados', ents.get('enviados', 0)],
                        ['Pendientes', ents.get('pendientes', 0)],
                        ['Tasa entrega a tiempo (%)', ents.get('tasa_entrega_tiempo_pct', 0)],
                    ])
    _autoajustar(ws_ents)

    ws_est = wb.create_sheet('Estudiantes')
    ws_est['A1'] = 'Avance por Estudiante'
    ws_est['A1'].font = Font(bold=True, color=ROJO)
    avances = datos.get('avance_por_estudiante', [])
    enc_est = ['ID', 'Nombre', 'Apellido', 'Código', 'Avance (%)', 'Nota (0-5)', 'Actividades con avance']
    filas_est = [
        [e.get('usuario_id'), e.get('nombre'), e.get('apellido'), e.get('codigo'),
         e.get('promedio_avance_pct'), e.get('nota_promedio_5'), e.get('actividades_con_avance')]
        for e in avances
    ]
    _escribir_tabla(ws_est, 3, enc_est, filas_est)
    _autoajustar(ws_est)

    ws_bajo = wb.create_sheet('Bajo rendimiento')
    ws_bajo['A1'] = f"Umbral: {datos.get('umbral_bajo_rendimiento', 3.0)}"
    ws_bajo['A1'].font = Font(bold=True, color=ROJO)
    bajo = datos.get('estudiantes_bajo_rendimiento', [])
    _escribir_tabla(ws_bajo, 3,
                    ['Nombre', 'Apellido', 'Código', 'Nota (0-5)', 'Avance (%)'],
                    [[e.get('nombre'), e.get('apellido'), e.get('codigo'),
                      e.get('nota_promedio_5'), e.get('promedio_avance_pct')]
                     for e in bajo])
    _autoajustar(ws_bajo)

    wb.save(ruta_destino)
    return ruta_destino


def generar_excel_estudiante(datos, ruta_destino):
    wb = Workbook()
    _hoja_info(wb)

    est = datos.get('estudiante', {})
    ws_info = wb.active
    ws_info['A4'] = 'Estudiante:'
    ws_info['B4'] = f"{est.get('nombre','')} {est.get('apellido','')}"
    ws_info['A5'] = 'Código:'
    ws_info['B5'] = est.get('codigo', '-')
    ws_info['A6'] = 'Proyecto:'
    ws_info['B6'] = datos.get('proyecto', {}).get('nombre', '-')

    ws_nota = wb.create_sheet('Nota Final')
    ws_nota['A1'] = 'Cálculo de Nota Final'
    ws_nota['A1'].font = Font(bold=True, color=ROJO)
    nota = datos.get('nota_final', {})
    comp_filas = [
        [comp.replace('_', ' ').title(), vals.get('nota', '-'), f"{vals.get('peso_aplicado','-')}%"]
        for comp, vals in nota.get('componentes', {}).items()
    ]
    comp_filas.append(['NOTA FINAL', nota.get('nota_final', '-'), ''])
    _escribir_tabla(ws_nota, 3, ['Componente', 'Nota (0-5)', 'Peso'], comp_filas)
    _autoajustar(ws_nota)

    ws_desemp = wb.create_sheet('Desempeño')
    d = datos.get('desempeno', {})
    _escribir_tabla(ws_desemp, 1,
                    ['Indicador', 'Valor'],
                    [
                        ['Actividades completadas', d.get('actividades_completadas', 0)],
                        ['Actividades vencidas', d.get('actividades_vencidas', 0)],
                        ['Total actividades equipo', d.get('total_actividades_equipo', 0)],
                        ['Avance promedio (%)', d.get('promedio_avance_pct', 0)],
                        ['Entregables aprobados', d.get('entregables_aprobados', 0)],
                        ['Entregables rechazados', d.get('entregables_rechazados', 0)],
                        ['Entregados a tiempo', d.get('entregables_a_tiempo', 0)],
                    ])
    _autoajustar(ws_desemp)

    ws_hist = wb.create_sheet('Historial Entregables')
    hist = datos.get('historial_entregables', [])
    _escribir_tabla(ws_hist, 1,
                    ['Título', 'Estado', 'Versión', 'Fecha envío', 'A tiempo', 'Retroalimentación'],
                    [
                        [e.get('titulo'), e.get('estado'), e.get('numero_version'),
                         e.get('fecha_envio'), 'Sí' if e.get('entregado_a_tiempo') else 'No',
                         e.get('retroalimentacion') or '']
                        for e in hist
                    ])
    _autoajustar(ws_hist)

    wb.save(ruta_destino)
    return ruta_destino


def generar_excel_indicadores(datos, ruta_destino):
    wb = Workbook()
    _hoja_info(wb)

    resumen_raw = datos.get('resumen_periodo', {})
    periodos_list = resumen_raw if isinstance(resumen_raw, list) else [resumen_raw]

    ws_info = wb.active
    ws_info['A4'] = 'Periodos exportados:'
    ws_info['B4'] = ', '.join(p.get('periodo_nombre', '-') for p in periodos_list)

    ws_res = wb.create_sheet('Resumen')
    _escribir_tabla(ws_res, 1,
                    ['Periodo', 'Total cursos', 'Cursos activos', 'Total estudiantes',
                     'Total proyectos', 'Proyectos activos', 'Avance prom. (%)', 'Aprobación (%)'],
                    [
                        [p.get('periodo_nombre', '-'),
                         p.get('total_cursos', 0),
                         p.get('cursos_activos', 0),
                         p.get('total_estudiantes', 0),
                         p.get('total_proyectos', 0),
                         p.get('proyectos_activos', 0),
                         p.get('avance_promedio_proyectos_pct', 0),
                         p.get('tasa_aprobacion_pct', 0)]
                        for p in periodos_list
                    ])
    _autoajustar(ws_res)

    ws_dist = wb.create_sheet('Distribución Notas')
    dist = datos.get('distribucion_notas', {})
    _escribir_tabla(ws_dist, 1,
                    ['Rango', 'Estudiantes'],
                    [
                        ['0.0 – 1.9', dist.get('rango_0_2', 0)],
                        ['2.0 – 2.9', dist.get('rango_2_3', 0)],
                        ['3.0 – 3.9', dist.get('rango_3_4', 0)],
                        ['4.0 – 5.0', dist.get('rango_4_5', 0)],
                        ['Promedio global', dist.get('nota_promedio_global', 0)],
                    ])
    _autoajustar(ws_dist)

    ws_proy = wb.create_sheet('Proyectos')
    proyectos = datos.get('proyectos', [])
    _escribir_tabla(ws_proy, 1,
                    ['Nombre', 'Curso', 'Estado', 'Avance (%)', 'Total entregables', 'Aprobados', 'Tasa a tiempo (%)'],
                    [
                        [p.get('nombre'), p.get('curso_nombre'), p.get('estado'),
                         p.get('porcentaje_progreso'), p.get('total_entregables'),
                         p.get('entregables_aprobados'), p.get('tasa_entrega_tiempo_pct')]
                        for p in proyectos
                    ])
    _autoajustar(ws_proy)

    ws_riesgo = wb.create_sheet('Estudiantes en Riesgo')
    riesgo = datos.get('estudiantes_riesgo_por_curso', [])
    _escribir_tabla(ws_riesgo, 1,
                    ['Curso', 'En riesgo', 'Total con avance'],
                    [[r.get('curso_nombre'), r.get('estudiantes_en_riesgo'), r.get('total_estudiantes_con_avance')]
                     for r in riesgo])
    _autoajustar(ws_riesgo)

    wb.save(ruta_destino)
    return ruta_destino
